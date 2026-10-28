#!/usr/bin/env python3
"""
India Pincode -> Vidhan Sabha mapper
====================================

Builds the hierarchy  Pincode -> Post Office / Village -> Vidhan Sabha -> Assembly No. -> District -> State
from three sets of KML boundary files:

  * LGD village boundaries          (e.g. lgd_villages.kml, lgd_villages (1).kml, ...)
  * Assembly constituency boundaries (e.g. LGD_Assembly_Constituencies.kml)
  * Pincode boundaries              (e.g. Datagov_Pincode_Boundaries.kml)

How the links are made
  * Every village is placed by an interior point of its boundary and tested against the
    pincode and constituency polygons (nearest polygon within ~2 km if it falls in a gap).
  * Every pincode is overlaid on the constituencies; the share of its area in each one is
    the "area overlap". The constituency with the largest overlap is the primary Vidhan Sabha.
  * Overlaps under 1% are dropped unless a village falls in that constituency.

Install once
    pip install lxml shapely pandas openpyxl

Build everything (put all the .kml files in one folder)
    python india_pincode_vidhan_sabha.py build --input-dir "C:/data/kml" --out output

    or name the files explicitly (wildcards are fine):
    python india_pincode_vidhan_sabha.py build --villages "kml/lgd_villages*.kml"
        --pincodes kml/Datagov_Pincode_Boundaries.kml --acs kml/LGD_Assembly_Constituencies.kml --out output
    (all on one line)

Outputs (in --out)
    pincode_vidhan_sabha.xlsx        Pincode Hierarchy, PIN-AC Overlaps, Vidhan Sabha summary
    village_hierarchy.csv            one row per village with its pincode, post office and Vidhan Sabha
    India_Pincode_Vidhan_Sabha.html  offline lookup page (open in any browser)
    lookup_data.json                 data used by the page and by the lookup command

Look up from the terminal (after build)
    python india_pincode_vidhan_sabha.py lookup 328025 --out output
    python india_pincode_vidhan_sabha.py lookup Daulatpur --out output
"""
import argparse
import datetime as dt
import glob
import json
import os
import sys
import time
from multiprocessing import Pool

import numpy as np
import pandas as pd
import shapely
from lxml import etree
from shapely.geometry import MultiPolygon, Polygon

try:
    sys.stdout.reconfigure(encoding="utf-8", errors="replace")
except Exception:
    pass

# Telangana's Lok Sabha seats: the AC file lists Telangana under ANDHRA PRADESH.
TELANGANA_PCS = {
    "ADILABAD(ST)", "BHONGIR", "CHEVELLA", "HYDERABAD", "KARIMNAGAR", "KHAMMAM", "MAHABUBABAD(ST)",
    "MAHBUBNAGAR", "MALKAJGIRI", "MEDAK", "NAGARKURNOOL(SC)", "NALGONDA", "NIZAMABAD",
    "PEDDAPALLE (SC)", "SECUNDERABAD", "WARANGAL(SC)", "ZAHIRABAD",
}
STATE_FIX = {"ORISSA": "ODISHA", "UTTARKHAND": "UTTARAKHAND"}
NEAREST_DEG = 0.02      # ~2 km: fallback search radius for points that fall in boundary gaps
MIN_OVERLAP_PCT = 1.0   # drop smaller pincode/constituency slivers unless a village is there


def log(msg):
    print(f"[{time.strftime('%H:%M:%S')}] {msg}", flush=True)


def txt(x):
    return "" if x is None or (isinstance(x, float) and x != x) else str(x).strip()


# ----------------------------------------------------------------------------- KML reading
def iter_placemarks(path):
    """Stream Placemarks from a (possibly huge) KML file without loading it all."""
    for _, pm in etree.iterparse(path, events=("end",), tag="{*}Placemark", huge_tree=True):
        yield pm
        pm.clear()
        while pm.getprevious() is not None:
            del pm.getparent()[0]


def placemark_attrs(pm):
    out = {}
    nm = pm.find("{*}name")
    out["_name"] = txt(nm.text) if nm is not None else ""
    for d in pm.iter("{*}Data"):
        v = d.find("{*}value")
        out[d.get("name")] = txt(v.text) if v is not None else ""
    for d in pm.iter("{*}SimpleData"):
        out[d.get("name")] = txt(d.text)
    return out


def _ring(el):
    c = el.find(".//{*}coordinates")
    if c is None or not c.text:
        return None
    pts = []
    for tok in c.text.split():
        p = tok.split(",")
        if len(p) >= 2:
            pts.append((float(p[0]), float(p[1])))
    return pts if len(pts) >= 4 else None


def placemark_geometry(pm):
    polys = []
    for pg in pm.iter("{*}Polygon"):
        ob = pg.find("{*}outerBoundaryIs")
        shell = _ring(ob) if ob is not None else None
        if not shell:
            continue
        holes = [h for h in (_ring(ib) for ib in pg.findall("{*}innerBoundaryIs")) if h]
        polys.append(Polygon(shell, holes))
    if not polys:
        return None
    g = polys[0] if len(polys) == 1 else MultiPolygon(polys)
    if not g.is_valid:
        g = shapely.make_valid(g)
    return g


def read_polygons(path):
    """Attributes + geometry for every Placemark (used for pincodes and constituencies)."""
    rows, geoms = [], []
    for pm in iter_placemarks(path):
        rows.append(placemark_attrs(pm))
        geoms.append(placemark_geometry(pm))
    return pd.DataFrame(rows), geoms


def read_villages(path):
    """Village attributes + one interior point per village (geometry is not kept, to save memory)."""
    rows = []
    for pm in iter_placemarks(path):
        a = placemark_attrs(pm)
        try:
            g = placemark_geometry(pm)
            p = g.representative_point() if g is not None and not g.is_empty else None
        except Exception:
            p = None
        rows.append({
            "village": a.get("_name") or a.get("vilname11") or a.get("vilnam_soi") or "(unnamed)",
            "state": a.get("stname", ""), "district": a.get("dtname", ""), "sub_district": a.get("sdtname", ""),
            "block": a.get("block_name", ""), "gram_panchayat": a.get("gp_name", ""),
            "village_lgd": a.get("vil_lgd", ""), "ac_no_in_file": a.get("ac_no", ""),
            "area_sq_km": round(float(a["shape_Area"]) / 1e6, 4) if a.get("shape_Area") else None,
            "lon": p.x if p is not None else np.nan, "lat": p.y if p is not None else np.nan,
            "source_file": os.path.basename(path),
        })
    return pd.DataFrame(rows)


# ----------------------------------------------------------------------------- spatial joins
def as_array(geoms):
    return np.array([g if g is not None else shapely.Point() for g in geoms], dtype=object)


def locate(points, polygons, rank=None):
    """Index of the polygon containing each point (-1 if none), plus how it was matched."""
    tree = shapely.STRtree(polygons)
    ok = ~shapely.is_empty(points)
    pi, gi = tree.query(points[ok], predicate="within")
    pi = np.nonzero(ok)[0][pi]
    m = pd.DataFrame({"p": pi, "g": gi})
    if rank is not None:                       # prefer e.g. current constituencies over old ones
        m["r"] = rank[m.g.values]
        m = m.sort_values(["p", "r"])
    m = m.drop_duplicates("p")
    idx = np.full(len(points), -1)
    idx[m.p.values] = m.g.values
    how = np.where(idx >= 0, "inside", "not found").astype(object)
    miss = np.nonzero((idx < 0) & ok)[0]
    if len(miss):
        mi, mg = tree.query_nearest(points[miss], max_distance=NEAREST_DEG, all_matches=False)
        idx[miss[mi]] = mg
        how[miss[mi]] = "nearest (<2 km)"
    return idx, how


def area_overlaps(pin_geoms, ac_geoms):
    """Share (%) of each pincode's area that lies in each constituency it touches."""
    tree = shapely.STRtree(ac_geoms)
    pi, ai = tree.query(pin_geoms, predicate="intersects")
    try:
        inter = shapely.intersection(pin_geoms[pi], ac_geoms[ai])
    except Exception:                           # rare GEOS topology error: fall back pair by pair
        inter = []
        for a, b in zip(pin_geoms[pi], ac_geoms[ai]):
            try:
                inter.append(a.intersection(b))
            except Exception:
                inter.append(a.buffer(0).intersection(b.buffer(0)))
        inter = np.array(inter, dtype=object)
    pin_area = shapely.area(pin_geoms)
    with np.errstate(divide="ignore", invalid="ignore"):
        pct = 100 * shapely.area(inter) / pin_area[pi]
    df = pd.DataFrame({"pin": pi, "ac": ai, "pct": pct})
    return df[df.pct > 0]


# ----------------------------------------------------------------------------- build
def find_inputs(args):
    def expand(pattern):
        return sorted(glob.glob(pattern)) if pattern else []
    villages, pincodes, acs = expand(args.villages), expand(args.pincodes), expand(args.acs)
    if args.input_dir:
        allk = sorted(glob.glob(os.path.join(args.input_dir, "*.kml")))
        villages = villages or [f for f in allk if "village" in os.path.basename(f).lower()]
        pincodes = pincodes or [f for f in allk if "pincode" in os.path.basename(f).lower()]
        acs = acs or [f for f in allk if "assembly" in os.path.basename(f).lower()]
    if not villages or len(pincodes) != 1 or len(acs) != 1:
        sys.exit(f"Could not find the inputs: {len(villages)} village file(s), {len(pincodes)} pincode file(s), "
                 f"{len(acs)} constituency file(s). Pass --input-dir or --villages/--pincodes/--acs.")
    return villages, pincodes[0], acs[0]


def build(args):
    t0 = time.time()
    vfiles, pin_file, ac_file = find_inputs(args)
    os.makedirs(args.out, exist_ok=True)
    log(f"{len(vfiles)} village file(s), pincodes: {os.path.basename(pin_file)}, constituencies: {os.path.basename(ac_file)}")

    # 1. read everything (villages in parallel worker processes)
    vfiles = sorted(vfiles, key=lambda f: -os.path.getsize(f))
    with Pool(max(1, min(args.workers, len(vfiles)))) as pool:
        job = pool.map_async(read_villages, vfiles)
        log("reading pincode boundaries ...")
        pins, pin_geoms = read_polygons(pin_file)
        log(f"  {len(pins):,} pincodes")
        log("reading constituency boundaries ...")
        acs, ac_geoms = read_polygons(ac_file)
        log(f"  {len(acs):,} constituency polygons")
        log("reading village boundaries (largest step) ...")
        v = pd.concat(job.get(), ignore_index=True)
    log(f"  {len(v):,} villages")
    pin_geoms, ac_geoms = as_array(pin_geoms), as_array(ac_geoms)

    # 2. clean constituency attributes
    acs["State"] = acs.get("st_name", "").map(txt).str.upper().replace(STATE_FIX)
    acs.loc[(acs.State == "ANDHRA PRADESH") & acs.get("pc_name", "").map(txt).isin(TELANGANA_PCS), "State"] = "TELANGANA"
    acs["Current"] = acs.get("status", "").map(txt).eq("")
    for c in ["ac_no", "ac_name", "pc_name", "dist_name"]:
        if c not in acs:
            acs[c] = ""
    acs["ac_no_int"] = pd.to_numeric(acs.ac_no, errors="coerce").astype("Int64")
    acs["pc_name"] = acs.pc_name.map(txt).str.replace("?", " - ", regex=False).str.replace("  ", " ").str.title()

    # 3. villages -> pincode and constituency
    log("placing villages in pincodes and constituencies ...")
    pts = shapely.points(v.lon.values, v.lat.values)
    v["pin_idx"], v["pin_match"] = locate(pts, pin_geoms)
    v["ac_idx"], v["ac_match"] = locate(pts, ac_geoms, rank=np.where(acs.Current, 0, 1))
    log(f"  inside a pincode: {(v.pin_match == 'inside').mean():.2%}   inside a constituency: {(v.ac_match == 'inside').mean():.2%}")

    # 4. pincode x constituency area overlap
    log("measuring pincode / constituency area overlap ...")
    ov = area_overlaps(pin_geoms, ac_geoms)
    log(f"  {len(ov):,} pincode-constituency overlaps")

    # 5. links = overlaps + villages per (pincode, constituency)
    vv = v[v.pin_idx >= 0].sort_values("village")
    vg = vv.groupby(["pin_idx", "ac_idx"]).agg(n=("village", "size"),
                                               names=("village", lambda s: "|".join(x.replace("|", "/") for x in s)))
    vg = vg.reset_index().rename(columns={"pin_idx": "pin", "ac_idx": "ac"})
    links = ov.merge(vg, on=["pin", "ac"], how="outer")
    links["pct"] = links.pct.fillna(0.0)
    links["n"] = links.n.fillna(0).astype(int)
    links["names"] = links.names.fillna("")
    links = links[(links.pct >= MIN_OVERLAP_PCT) | (links.n > 0)]
    links = links.sort_values(["pin", "pct", "n"], ascending=[True, False, False]).reset_index(drop=True)

    # 6. tables
    mode = lambda s: s.mode().iat[0] if s.notna().any() and len(s.mode()) else ""
    pin_vil = vv.groupby("pin_idx").agg(villages=("village", "size"), district=("district", mode), state=("state", mode))
    ac_links = links[links.ac >= 0]
    primary = ac_links.drop_duplicates("pin").set_index("pin")
    most_vil = ac_links[ac_links.n > 0].sort_values(["pin", "n"], ascending=[True, False]).drop_duplicates("pin").set_index("pin")
    others = ac_links.groupby("pin").apply(
        lambda d: "; ".join(f"{txt(acs.ac_name.iat[a])} ({txt(acs.ac_no.iat[a])}, {p:.1f}%)"
                            for a, p in zip(d.ac.iloc[1:], d.pct.iloc[1:])), include_groups=False)

    def ac_field(idx_series, col):
        return [txt(acs[col].iat[i]) if i >= 0 else "" for i in idx_series]

    P = pd.DataFrame({"Pincode": pins.get("Pincode", "").map(txt), "Post Office": pins.get("Office_Name", "").map(txt)})
    k = np.arange(len(P))
    P["Villages"] = pd.Series(k).map(pin_vil.villages).fillna(0).astype(int).values
    prim_ac = pd.Series(k).map(primary.ac).fillna(-1).astype(int).values
    P["Vidhan Sabha"] = ac_field(prim_ac, "ac_name")
    P["Assembly No."] = [acs.ac_no_int.iat[i] if i >= 0 else pd.NA for i in prim_ac]
    P["Area Overlap %"] = pd.Series(k).map(primary.pct).round(2).values
    P["District"] = [pin_vil.district.get(i) or (txt(acs.dist_name.iat[a]).title() if a >= 0 else "") for i, a in zip(k, prim_ac)]
    P["State"] = [pin_vil.state.get(i) or (acs.State.iat[a] if a >= 0 else "") for i, a in zip(k, prim_ac)]
    P["Other Vidhan Sabhas"] = pd.Series(k).map(others).fillna("").values
    mv = pd.Series(k).map(most_vil.ac).fillna(-1).astype(int).values
    P["Vidhan Sabha (most villages)"] = ac_field(mv, "ac_name")
    P["Lok Sabha"] = ac_field(prim_ac, "pc_name")
    P["AC Status"] = ["Current" if a >= 0 and acs.Current.iat[a] else ("Pre-delimitation" if a >= 0 else "") for a in prim_ac]
    for src, dst in [("Division", "Division"), ("Region", "Region"), ("Circle", "Circle")]:
        P[dst] = pins.get(src, "").map(txt) if src in pins else ""

    V = pd.DataFrame({
        "Pincode": [P.Pincode.iat[i] if i >= 0 else "" for i in v.pin_idx],
        "Post Office": [P["Post Office"].iat[i] if i >= 0 else "" for i in v.pin_idx],
        "Village": v.village, "Vidhan Sabha": ac_field(v.ac_idx, "ac_name"),
        "Assembly No.": [acs.ac_no_int.iat[i] if i >= 0 else pd.NA for i in v.ac_idx],
        "District": v.district, "State": v.state, "Village LGD": v.village_lgd, "Gram Panchayat": v.gram_panchayat,
        "Block": v.block, "Sub-District": v.sub_district, "Lok Sabha": ac_field(v.ac_idx, "pc_name"),
        "AC No. in village file": v.ac_no_in_file, "Area (sq km)": v.area_sq_km,
        "Latitude": v.lat.round(6), "Longitude": v.lon.round(6),
        "Pincode Match": v.pin_match, "AC Match": v.ac_match, "Source File": v.source_file,
    }).sort_values(["State", "District", "Assembly No.", "Pincode", "Village"], na_position="last")

    O = pd.DataFrame({
        "Pincode": P.Pincode.values[ac_links.pin.values], "Post Office": P["Post Office"].values[ac_links.pin.values],
        "Rank": ac_links.groupby("pin").cumcount().values + 1,
        "Vidhan Sabha": ac_field(ac_links.ac.values, "ac_name"),
        "Assembly No.": acs.ac_no_int.values[ac_links.ac.values],
        "Area Overlap %": ac_links.pct.round(2).values, "Villages": ac_links.n.values,
        "Lok Sabha": ac_field(ac_links.ac.values, "pc_name"), "State (AC file)": acs.State.values[ac_links.ac.values],
    })
    A = O.groupby(["State (AC file)", "Assembly No.", "Vidhan Sabha"], dropna=False).agg(
        Pincodes=("Pincode", "nunique"), Villages=("Villages", "sum"),
        Pincode_List=("Pincode", lambda s: ", ".join(sorted(set(s))))).reset_index().rename(columns={"Pincode_List": "Pincode List"})

    # 7. write outputs
    xlsx = os.path.join(args.out, "pincode_vidhan_sabha.xlsx")
    log(f"writing {xlsx} ...")
    P_sorted = (P.assign(_blank=P.State.eq(""))
                .sort_values(["_blank", "State", "District", "Assembly No.", "Pincode"], na_position="last")
                .drop(columns="_blank"))
    write_excel(xlsx, {"Pincode Hierarchy": P_sorted,
                       "PIN-AC Overlaps": O, "Vidhan Sabha": A})
    vcsv = os.path.join(args.out, "village_hierarchy.csv")
    log(f"writing {vcsv} ...")
    V.to_csv(vcsv, index=False, encoding="utf-8-sig")

    data = lookup_payload(P, acs, links, v)
    with open(os.path.join(args.out, "lookup_data.json"), "w", encoding="utf-8") as f:
        json.dump(data, f, separators=(",", ":"), ensure_ascii=False)
    html = os.path.join(args.out, "India_Pincode_Vidhan_Sabha.html")
    log(f"writing {html} ...")
    payload = json.dumps(data, separators=(",", ":"), ensure_ascii=False).replace("</", "<\\/")
    page = HTML_TEMPLATE.replace("__DATA__", payload).replace("__GEN__", dt.date.today().strftime("%d %B %Y"))
    with open(html, "w", encoding="utf-8") as f:
        f.write(page)

    multi = (ac_links[ac_links.pct >= 5].groupby("pin").size() > 1).sum()
    log(f"done in {time.time() - t0:,.0f}s -> {len(P):,} pincodes, {len(V):,} villages, "
        f"{len(O):,} pincode-constituency links, {multi:,} pincodes in 2+ constituencies (>=5% each)")


def write_excel(path, sheets):
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter
    with pd.ExcelWriter(path, engine="openpyxl") as xw:
        for name, df in sheets.items():
            df.to_excel(xw, sheet_name=name, index=False)
            ws = xw.sheets[name]
            for j, col in enumerate(df.columns, 1):
                c = ws.cell(1, j)
                c.font = Font(name="Arial", bold=True, color="FFFFFF")
                c.fill = PatternFill("solid", fgColor="1F4E5A")
                c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
                width = max([len(str(col))] + [len(str(x)) for x in df[col].head(500)]) + 2
                ws.column_dimensions[get_column_letter(j)].width = min(max(width, 9), 45)
            ws.freeze_panes = "B2"
            ws.auto_filter.ref = ws.dimensions


def lookup_payload(P, acs, links, v):
    """Compact structure shared by the HTML page and the lookup command."""
    states = sorted(set(acs.State) | set(v.state.map(txt).str.upper()) | set(P.State.map(txt).str.upper()) - {""})
    st_i = {s: i for i, s in enumerate(states)}
    used = sorted(set(links.ac[links.ac >= 0]))
    ac_i = {a: i for i, a in enumerate(used)}
    ac_rows = [[st_i.get(acs.State.iat[a], -1), (int(acs.ac_no_int.iat[a]) if pd.notna(acs.ac_no_int.iat[a]) else None),
                txt(acs.ac_name.iat[a]), txt(acs.pc_name.iat[a]), 1 if acs.Current.iat[a] else 0,
                txt(acs.dist_name.iat[a]).title()] for a in used]
    districts, d_i = [], {}
    pins = []
    for r in P.itertuples(index=False):
        d = txt(r.District)
        if d and d not in d_i:
            d_i[d] = len(districts)
            districts.append(d)
        pins.append([r.Pincode, r[1], txt(r.Division), txt(r.Region), txt(r.Circle),
                     st_i.get(txt(r.State).upper(), -1), d_i.get(d, -1), int(r.Villages)])
    link_rows = [[int(r.pin), ac_i.get(int(r.ac), -1), round(float(r.pct), 2), int(r.n), r.names]
                 for r in links.itertuples(index=False)]
    return {"states": [s.title() for s in states], "districts": districts, "acs": ac_rows, "pins": pins,
            "links": link_rows, "generated": dt.date.today().isoformat()}


# ----------------------------------------------------------------------------- lookup
def lookup(args):
    path = os.path.join(args.out, "lookup_data.json")
    if not os.path.exists(path):
        sys.exit(f"{path} not found. Run the build command first.")
    with open(path, encoding="utf-8") as f:
        D = json.load(f)
    pins, links, acs, states, dists = D["pins"], D["links"], D["acs"], D["states"], D["districts"]
    by_pin = {}
    for l in links:
        by_pin.setdefault(l[0], []).append(l)
    ac_name = lambda a: acs[a][2] if a >= 0 else "(no constituency)"
    ac_no = lambda a: acs[a][1] if a >= 0 else "-"
    q = args.query.strip()

    if q.isdigit():
        hit = [i for i, p in enumerate(pins) if p[0] == q]
        if not hit:
            sys.exit(f"PIN code {q} is not in the data.")
        i = hit[0]
        p = pins[i]
        ls = by_pin.get(i, [])
        top = next((l for l in ls if l[1] >= 0), None)
        print(f"\nPIN {p[0]}  |  {p[1]}")
        print("-" * 60)
        print(f"Pincode      : {p[0]}")
        print(f"Post Office  : {p[1]}")
        print(f"Villages     : {sum(l[3] for l in ls):,}")
        print(f"Vidhan Sabha : {ac_name(top[1]) if top else '-'}")
        print(f"Assembly No. : {ac_no(top[1]) if top else '-'}")
        print(f"District     : {dists[p[6]] if p[6] >= 0 else '-'}")
        print(f"State        : {states[p[5]] if p[5] >= 0 else '-'}")
        print("\nConstituencies by area overlap:")
        for k, l in enumerate([l for l in ls if l[1] >= 0], 1):
            a = acs[l[1]]
            print(f"  {k}. {a[2]} (AC {a[1]}) - {l[2]:.2f}% of area, {l[3]} villages, Lok Sabha: {a[3]}, "
                  f"{'current' if a[4] else 'pre-delimitation'}")
            if args.villages and l[4]:
                print("     " + ", ".join(l[4].split("|")))
        return

    s = q.lower()
    found = []
    for l in links:
        if not l[4]:
            continue
        for name in l[4].split("|"):
            if s in name.lower():
                found.append((not name.lower().startswith(s), name, l))
    found.sort(key=lambda x: (x[0], x[1]))
    if not found:
        sys.exit(f"No village matching '{q}'.")
    print(f"\n{len(found):,} village(s) matching '{q}'" + (" (first 50 shown)" if len(found) > 50 else ""))
    print(f"{'Village':<28}{'PIN':<8}{'Post office':<26}{'Vidhan Sabha':<26}{'AC':>4}  District / State")
    for _, name, l in found[:50]:
        p = pins[l[0]]
        print(f"{name[:27]:<28}{p[0]:<8}{p[1][:25]:<26}{ac_name(l[1])[:25]:<26}{str(ac_no(l[1])):>4}  "
              f"{dists[p[6]] if p[6] >= 0 else '-'} / {states[p[5]] if p[5] >= 0 else '-'}")


# ----------------------------------------------------------------------------- CLI
def main():
    ap = argparse.ArgumentParser(description="India Pincode -> Vidhan Sabha mapper", formatter_class=argparse.RawDescriptionHelpFormatter,
                                 epilog=__doc__.split("Install once")[1] if "Install once" in __doc__ else None)
    sub = ap.add_subparsers(dest="cmd", required=True)
    b = sub.add_parser("build", help="read the KML files and write Excel, CSV, JSON and the HTML lookup page")
    b.add_argument("--input-dir", help="folder holding all the .kml files (picked by name: village / pincode / assembly)")
    b.add_argument("--villages", help='village KML file(s), wildcards allowed, e.g. "kml/lgd_villages*.kml"')
    b.add_argument("--pincodes", help="pincode boundary KML file")
    b.add_argument("--acs", help="assembly constituency KML file")
    b.add_argument("--out", default="output", help="output folder (default: output)")
    b.add_argument("--workers", type=int, default=max(1, (os.cpu_count() or 2) - 1), help="parallel processes for village files")
    l = sub.add_parser("lookup", help="look up a PIN code or village name from a previous build")
    l.add_argument("query", help="6-digit PIN code or part of a village name")
    l.add_argument("--out", default="output", help="folder used for build (default: output)")
    l.add_argument("--villages", action="store_true", help="also list the village names for a PIN code")
    args = ap.parse_args()
    build(args) if args.cmd == "build" else lookup(args)


HTML_TEMPLATE = r'''<!doctype html>
<html lang="en">
<head>
<meta charset="utf-8">
<meta name="viewport" content="width=device-width,initial-scale=1">
<title>India Pincode → Vidhan Sabha</title>
<link rel="preconnect" href="https://fonts.googleapis.com">
<link rel="stylesheet" href="https://fonts.googleapis.com/css2?family=Inter:wght@400;500;600;700;800&display=swap">
<style>
:root{--bg:#f4f7fb;--surface:#fff;--ink:#111827;--muted:#64748b;--line:#e2e8f0;--navy:#0b1736;--blue:#2563eb;--blue2:#4f46e5;--green:#047857;--greenbg:#ecfdf5;--amber:#c2410c;--amberbg:#fff7ed;--shadow:0 12px 35px rgba(15,23,42,.08)}
*{box-sizing:border-box}
body{margin:0;font-family:Inter,system-ui,-apple-system,"Segoe UI",sans-serif;color:var(--ink);background:linear-gradient(180deg,#f8fbff,#eef3f9);min-height:100vh}
header{height:72px;background:#fffffff0;backdrop-filter:blur(16px);border-bottom:1px solid var(--line);display:flex;align-items:center;justify-content:space-between;gap:12px;padding:0 34px;position:sticky;top:0;z-index:20}
.logo{display:flex;align-items:center;gap:11px;font-weight:800;font-size:17px;min-width:0}
.logoMark{width:34px;height:34px;border-radius:10px;background:linear-gradient(135deg,#2563eb,#4f46e5);display:grid;place-items:center;color:#fff;flex:none}
.logoMark svg{width:19px;height:19px}
.badge{font-size:10px;font-weight:800;background:#eff6ff;color:#1d4ed8;border:1px solid #bfdbfe;padding:5px 9px;border-radius:999px;letter-spacing:.04em;white-space:nowrap}
.hdrlink{font-size:13px;font-weight:600;color:var(--muted);text-decoration:none}
.hdrlink:hover{color:var(--blue)}
.wrap{max-width:1240px;margin:auto;padding:40px 24px 80px}
.hero{position:relative;overflow:hidden;border-radius:28px;padding:58px 28px 50px;text-align:center;background:radial-gradient(circle at 15% 20%,#3b82f647,transparent 30%),radial-gradient(circle at 85% 15%,#06b6d433,transparent 28%),linear-gradient(135deg,#0b1736,#172554 58%,#1e3a8a);box-shadow:0 20px 55px #2563eb24;color:#fff}
.eyebrow{display:inline-flex;gap:6px;align-items:center;background:#ffffff1a;border:1px solid #ffffff26;border-radius:999px;padding:7px 12px;font-size:11px;font-weight:700;letter-spacing:.08em;text-transform:uppercase;color:#dbeafe}
.eyebrow i{width:7px;height:7px;border-radius:50%;background:#34d399;box-shadow:0 0 0 4px #34d39933}
.hero h1{font-size:46px;line-height:1.06;letter-spacing:-1.6px;margin:18px auto 14px;max-width:860px;text-wrap:balance}
.hero p{color:#cbd5e1;font-size:16px;line-height:1.65;margin:0 auto 26px;max-width:740px}
.search{display:flex;gap:9px;max-width:780px;margin:auto;background:#ffffff1a;padding:7px;border-radius:17px;border:1px solid #ffffff24}
.search input{flex:1;min-width:0;border:0;border-radius:11px;padding:16px 18px;font:inherit;font-size:16px;outline:none;color:var(--ink)}
.search input:focus{box-shadow:0 0 0 3px #60a5fa88}
.search button{border:0;background:linear-gradient(135deg,#60a5fa,#2563eb);color:#fff;border-radius:11px;padding:0 26px;font:inherit;font-weight:700;cursor:pointer;white-space:nowrap}
.search button:hover{filter:brightness(1.08)}
.examples{font-size:12px;color:#bfdbfe;margin-top:14px;display:flex;flex-wrap:wrap;gap:6px;justify-content:center;align-items:center}
.examples button{font:inherit;font-size:12px;font-weight:700;color:#fff;background:#ffffff14;border:1px solid #ffffff26;border-radius:999px;padding:4px 10px;cursor:pointer}
.examples button:hover{background:#ffffff26}
.cards{display:grid;grid-template-columns:repeat(6,1fr);gap:14px;margin:22px 0}
.card{background:#fffffff2;border:1px solid var(--line);border-radius:17px;padding:17px;box-shadow:0 5px 20px #0f172a0a;display:flex;align-items:center;gap:12px;min-width:0}
.icon{width:40px;height:40px;border-radius:12px;display:grid;place-items:center;background:#eff6ff;color:#2563eb;flex:none}
.icon svg{width:20px;height:20px}
.card .n{font-size:22px;font-weight:800;font-variant-numeric:tabular-nums;letter-spacing:-.3px}
.card .l{font-size:12px;color:var(--muted)}
.panel{background:#fffffff2;border:1px solid var(--line);border-radius:20px;overflow:hidden;box-shadow:var(--shadow);margin-bottom:22px}
.panelhead{display:flex;justify-content:space-between;align-items:center;gap:12px;flex-wrap:wrap;padding:16px 20px;border-bottom:1px solid var(--line);background:#fbfdff}
.panelhead h2{margin:0;font-size:15px;font-weight:700}
.panelhead .sub{font-size:12px;color:var(--muted)}
.toolbar{padding:15px 20px;display:flex;gap:10px;flex-wrap:wrap;border-bottom:1px solid var(--line);background:#fbfdff;align-items:center}
.toolbar select,.toolbar input{border:1px solid #dbe2ea;background:#fff;padding:10px 12px;border-radius:10px;outline:none;font:inherit;font-size:13px;color:var(--ink);max-width:100%}
.toolbar input{flex:1 1 240px}
.toolbar select:focus,.toolbar input:focus{border-color:#93c5fd;box-shadow:0 0 0 3px #bfdbfe}
.btn{border:1px solid #dbe2ea;background:#fff;padding:10px 14px;border-radius:10px;font:inherit;font-size:13px;font-weight:600;color:var(--ink);cursor:pointer;white-space:nowrap}
.btn:hover{border-color:#93c5fd;color:var(--blue)}
.btn:focus-visible,.examples button:focus-visible,.search button:focus-visible,.linkbtn:focus-visible{outline:2px solid #60a5fa;outline-offset:2px}
.result{padding:25px}
.empty{text-align:center;color:var(--muted);padding:50px 20px;font-size:14px;line-height:1.6}
.resulthead{display:flex;justify-content:space-between;gap:20px;align-items:flex-start;margin-bottom:18px}
.pin{font-size:34px;font-weight:800;letter-spacing:-.5px;font-variant-numeric:tabular-nums}
.meta{color:var(--muted);font-size:13px;margin-top:5px}
.pills{display:flex;gap:6px;flex-wrap:wrap;justify-content:flex-end}
.pill{font-size:11px;font-weight:700;border-radius:999px;padding:5px 9px;background:#f1f5f9;color:#475569;white-space:nowrap}
.pill.primary{background:var(--greenbg);color:var(--green)}
.pill.warning{background:var(--amberbg);color:var(--amber)}
.pill.info{background:#eff6ff;color:#1d4ed8}
.path{display:flex;flex-wrap:wrap;align-items:stretch;gap:6px;margin:0 0 20px;padding:14px;border:1px solid #dbeafe;border-radius:15px;background:linear-gradient(180deg,#f8fbff,#f1f6ff)}
.step{display:grid;gap:2px;padding:8px 12px;border-radius:11px;background:#fff;border:1px solid var(--line);min-width:0}
.step small{font-size:10px;font-weight:700;letter-spacing:.07em;text-transform:uppercase;color:var(--muted)}
.step span{font-size:14px;font-weight:700;overflow-wrap:anywhere}
.arrow{align-self:center;color:#93a3b8;font-weight:700}
.grid{display:grid;grid-template-columns:repeat(auto-fit,minmax(min(100%,285px),1fr));gap:13px}
.ac{border:1px solid #e2e8f0;border-radius:15px;padding:18px;background:linear-gradient(180deg,#fff,#f9fbff);display:grid;gap:2px;align-content:start}
.ac.top{border-color:#a7f3d0;box-shadow:0 0 0 3px #ecfdf5}
.ac h3{margin:0 0 8px;font-size:17px;display:flex;justify-content:space-between;gap:10px;align-items:flex-start}
.ac .row{display:flex;justify-content:space-between;gap:12px;font-size:13px;padding:5px 0}
.ac .row b{text-align:right}
.label{color:var(--muted)}
.bar{height:7px;border-radius:99px;background:#e2e8f0;overflow:hidden;margin:2px 0 6px}
.bar i{display:block;height:100%;border-radius:99px;background:linear-gradient(90deg,#60a5fa,#2563eb)}
.ac.top .bar i{background:linear-gradient(90deg,#34d399,#059669)}
.villages{margin-top:22px;display:grid;gap:14px}
.villages h4{margin:0 0 8px;font-size:13px;font-weight:700;display:flex;gap:8px;align-items:center;flex-wrap:wrap}
.chips{display:flex;flex-wrap:wrap;gap:6px}
.chip{font-size:12px;padding:4px 9px;border-radius:8px;background:#f1f5f9;color:#334155;border:1px solid #e2e8f0}
.chip.hit{background:#fef9c3;border-color:#fde047;color:#713f12;font-weight:700}
.linkbtn{font:inherit;font-size:12px;font-weight:700;color:var(--blue);background:none;border:0;padding:4px 2px;cursor:pointer}
.extra{margin-top:18px;font-size:12px;color:var(--muted);display:flex;gap:16px;flex-wrap:wrap}
.extra b{color:var(--ink);font-weight:600}
.tablewrap{overflow:auto;max-height:640px}
.table{width:100%;border-collapse:collapse;font-size:13px;font-variant-numeric:tabular-nums}
.table th,.table td{padding:11px 14px;border-bottom:1px solid #eef2f7;text-align:left;white-space:nowrap}
.table th{background:#f8fafc;font-size:10px;text-transform:uppercase;letter-spacing:.08em;color:#64748b;position:sticky;top:0;z-index:1}
.table td.num,.table th.num{text-align:right}
.table tbody tr:hover td{background:#f8fbff}
.table a{color:#2563eb;font-weight:700;text-decoration:none}
.table a:hover{text-decoration:underline}
.tablenote{padding:12px 20px;font-size:12px;color:var(--muted);border-top:1px solid var(--line);background:#fbfdff}
.mini{display:inline-block;width:46px;height:6px;border-radius:9px;background:#e2e8f0;vertical-align:middle;margin-right:6px;overflow:hidden}
.mini i{display:block;height:100%;background:#2563eb}
footer{color:#94a3b8;text-align:center;font-size:12px;line-height:1.7;padding-top:10px;max-width:880px;margin:auto}
mark{background:#fef08a;color:inherit;border-radius:3px;padding:0 1px}
@media(max-width:1100px){.cards{grid-template-columns:repeat(3,1fr)}}
@media(max-width:900px){.hero h1{font-size:38px}}
@media(max-width:650px){header{padding:0 16px}.hdrlink{display:none}.wrap{padding:22px 14px 55px}.hero{padding:40px 17px 36px;border-radius:22px}.hero h1{font-size:31px;letter-spacing:-1px}.search{flex-direction:column}.search button{height:48px}.cards{grid-template-columns:1fr 1fr}.resulthead{flex-direction:column}.pills{justify-content:flex-start}.result{padding:18px}.badge{display:none}}
@media(prefers-reduced-motion:no-preference){.ac,.step{animation:rise .35s ease-out both}@keyframes rise{from{opacity:.001;transform:translateY(6px)}to{opacity:1;transform:none}}}
</style>
</head>
<body>
<header>
  <div class="logo"><div class="logoMark"><svg viewBox="0 0 24 24" fill="none" stroke="currentColor" stroke-width="2.2" stroke-linecap="round" stroke-linejoin="round"><path d="M12 21s-6.5-5.8-6.5-11A6.5 6.5 0 0 1 18.5 10c0 5.2-6.5 11-6.5 11z"/><circle cx="12" cy="10" r="2.4"/></svg></div>India Pincode → Vidhan Sabha <span class="badge">ALL INDIA · VILLAGE LEVEL</span></div>
  <a class="hdrlink" href="#explorer">Browse all PIN codes ↓</a>
</header>
<main class="wrap">
<section class="hero">
  <div class="eyebrow"><i></i>Boundary + village aware lookup</div>
  <h1>Find the Vidhan Sabha and villages behind any Indian PIN code.</h1>
  <p>Search a six-digit PIN to see every Assembly Constituency it overlaps, ranked by share of area, with the villages that fall in each. Or type a village name to find its PIN code and constituency.</p>
  <div class="search"><input id="q" placeholder="Enter a 6-digit PIN code or a village name, e.g. 328025 or Rajakhera" autocomplete="off" aria-label="PIN code or village name"><button id="go" type="button">Search</button></div>
  <div class="examples">Try: <button type="button" data-ex="411041">411041</button><button type="button" data-ex="403503">403503</button><button type="button" data-ex="328025">328025</button><button type="button" data-ex="110001">110001</button><button type="button" data-ex="Daulatpur">Daulatpur</button></div>
</section>

<section class="cards" aria-label="Dataset totals">
  <div class="card"><div class="icon"><svg viewBox="0 0 24 24" fill="none" stroke="currentColor" stroke-width="2" stroke-linecap="round"><path d="M12 21s-6.5-5.8-6.5-11A6.5 6.5 0 0 1 18.5 10c0 5.2-6.5 11-6.5 11z"/><circle cx="12" cy="10" r="2.2"/></svg></div><div><div class="n" id="kPins">—</div><div class="l">PIN codes</div></div></div>
  <div class="card"><div class="icon"><svg viewBox="0 0 24 24" fill="none" stroke="currentColor" stroke-width="2" stroke-linecap="round" stroke-linejoin="round"><path d="M3 21h18M5 21V10l7-5 7 5v11M9 21v-6h6v6"/></svg></div><div><div class="n" id="kVil">—</div><div class="l">Villages mapped</div></div></div>
  <div class="card"><div class="icon"><svg viewBox="0 0 24 24" fill="none" stroke="currentColor" stroke-width="2" stroke-linecap="round" stroke-linejoin="round"><path d="M4 10h16M6 10v8M10 10v8M14 10v8M18 10v8M3 21h18M12 3l9 5H3z"/></svg></div><div><div class="n" id="kAc">—</div><div class="l">Vidhan Sabhas</div></div></div>
  <div class="card"><div class="icon"><svg viewBox="0 0 24 24" fill="none" stroke="currentColor" stroke-width="2" stroke-linecap="round" stroke-linejoin="round"><path d="M3 6l6-3 6 3 6-3v15l-6 3-6-3-6 3z"/><path d="M9 3v15M15 6v15"/></svg></div><div><div class="n" id="kDist">—</div><div class="l">Districts</div></div></div>
  <div class="card"><div class="icon"><svg viewBox="0 0 24 24" fill="none" stroke="currentColor" stroke-width="2" stroke-linecap="round" stroke-linejoin="round"><path d="M4 4h16v16H4z"/><path d="M4 12h16M12 4v16"/></svg></div><div><div class="n" id="kState">—</div><div class="l">States / UTs</div></div></div>
  <div class="card"><div class="icon"><svg viewBox="0 0 24 24" fill="none" stroke="currentColor" stroke-width="2" stroke-linecap="round" stroke-linejoin="round"><circle cx="9" cy="12" r="6"/><circle cx="15" cy="12" r="6"/></svg></div><div><div class="n" id="kMulti">—</div><div class="l">PINs in 2+ constituencies</div></div></div>
</section>

<section class="panel" aria-live="polite">
  <div id="result" class="result"><div class="empty">Enter a PIN code or village name above to see its full hierarchy:<br><b>Pincode → Post office / Villages → Vidhan Sabha → Assembly No. → District → State</b></div></div>
</section>

<section class="panel" id="explorer">
  <div class="panelhead"><h2>All PIN codes</h2><span class="sub" id="count"></span></div>
  <div class="toolbar">
    <select id="stateFilter" aria-label="State"><option value="">All states / UTs</option></select>
    <select id="distFilter" aria-label="District" disabled><option value="">All districts</option></select>
    <input id="tableSearch" placeholder="Filter by PIN, office, district or constituency" aria-label="Filter table">
    <button class="btn" id="csv" type="button">Download CSV</button>
  </div>
  <div class="tablewrap">
    <table class="table"><thead><tr><th>PIN</th><th>Post office</th><th class="num">Villages</th><th>Vidhan Sabha</th><th class="num">AC No.</th><th>District</th><th>State</th><th class="num">Overlap</th><th class="num">ACs</th></tr></thead>
    <tbody id="tbody"></tbody></table>
  </div>
  <div class="tablenote" id="tablenote"></div>
</section>

<footer>
  Built from the LGD village boundaries (12 KML files), LGD_Assembly_Constituencies.kml and Datagov_Pincode_Boundaries.kml. Overlap is the share of each PIN code's area inside a constituency; constituencies under 1% are left out unless a village falls in them. Villages are placed by an interior point of their boundary. Some states' constituency boundaries are pre-delimitation in the source file and are marked as such. Generated __GEN__.
</footer>
</main>

<script>
const D=__DATA__;
(function(){
const PINS=D.pins,LINKS=D.links,ACS=D.acs,ST=D.states,DI=D.districts;
const NF=new Intl.NumberFormat('en-IN');
const $=id=>document.getElementById(id);
const esc=x=>String(x??'').replace(/[&<>"']/g,m=>({'&':'&amp;','<':'&lt;','>':'&gt;','"':'&quot;',"'":'&#39;'}[m]));
// pin -> link range (links are sorted by pin)
const start=new Int32Array(PINS.length).fill(-1),end=new Int32Array(PINS.length).fill(-1);
LINKS.forEach((l,i)=>{if(start[l[0]]<0)start[l[0]]=i;end[l[0]]=i+1});
const pinIdx=Object.create(null);PINS.forEach((p,i)=>pinIdx[p[0]]=i);
const linksOf=i=>start[i]<0?[]:LINKS.slice(start[i],end[i]);
const acName=a=>a<0?'No constituency':(ACS[a][2]||'Unnamed');
const acNo=a=>a<0?'—':(ACS[a][1]??'—');
const stateOf=p=>p[5]<0?'':ST[p[5]];
const distOf=p=>p[6]<0?'':DI[p[6]];
// per-pin summary for the table
const SUM=PINS.map((p,i)=>{const ls=linksOf(i).filter(l=>l[1]>=0);const top=ls[0];return{i,ac:top?top[1]:-1,pct:top?top[2]:null,n:ls.length,second:ls[1]?ls[1][2]:0}});

// ---- totals
const vilTotal=LINKS.reduce((a,l)=>a+l[3],0);
$('kPins').textContent=NF.format(PINS.length);
$('kVil').textContent=NF.format(vilTotal);
$('kAc').textContent=NF.format(new Set(ACS.map(a=>a[0]+'|'+a[1]+'|'+a[2])).size);
$('kDist').textContent=NF.format(new Set(PINS.filter(p=>p[7]>0).map(p=>p[6])).size);
$('kState').textContent=NF.format(new Set(PINS.map(p=>p[5]).filter(x=>x>=0)).size);
$('kMulti').textContent=NF.format(SUM.filter(s=>s.second>=5).length);

// ---- village index (built on first name search)
let VN=null,VL=null,VNAME=null;
function buildVillageIndex(){
  VN=[];VL=[];VNAME=[];
  LINKS.forEach((l,li)=>{if(!l[4])return;for(const n of l[4].split('|')){VNAME.push(n);VN.push(n.toLowerCase());VL.push(li)}});
}

// ---- render a PIN
let lastHit=null;
function showPin(pin,hitVillage){
  const i=pinIdx[pin];
  const R=$('result');
  if(i===undefined){R.innerHTML=`<div class="empty">PIN code <b>${esc(pin)}</b> is not in this dataset.</div>`;return}
  const p=PINS[i],ls=linksOf(i);
  const acLinks=ls.filter(l=>l[1]>=0),noAc=ls.filter(l=>l[1]<0);
  const top=acLinks[0];
  const nVil=ls.reduce((a,l)=>a+l[3],0);
  const step=(k,v)=>`<div class="step"><small>${k}</small><span>${esc(v||'—')}</span></div>`;
  const arrow='<span class="arrow" aria-hidden="true">→</span>';
  const path=[step('Pincode',p[0]),step('Post office',p[1]),step('Villages',NF.format(nVil)),step('Vidhan Sabha',top?acName(top[1]):'—'),step('Assembly No.',top?acNo(top[1]):'—'),step('District',distOf(p)),step('State',stateOf(p))].join(arrow);
  const cards=acLinks.map((l,k)=>{const a=ACS[l[1]];return `<div class="ac${k===0?' top':''}"><h3><span>${esc(acName(l[1]))}</span><span class="pill ${k===0?'primary':''}">${k===0?'Primary':'#'+(k+1)}</span></h3>
    <div class="row"><span class="label">Assembly No.</span><b>${esc(acNo(l[1]))}</b></div>
    <div class="row"><span class="label">Area overlap</span><b>${l[2].toFixed(2)}%</b></div><div class="bar"><i style="width:${Math.max(1,Math.min(100,l[2]))}%"></i></div>
    <div class="row"><span class="label">Villages in this part</span><b>${NF.format(l[3])}</b></div>
    <div class="row"><span class="label">Lok Sabha</span><b>${esc(a[3]||'—')}</b></div>
    <div class="row"><span class="label">State</span><b>${esc(ST[a[0]])}</b></div>
    <div class="row"><span class="label">Boundary</span><span class="pill ${a[4]?'info':'warning'}">${a[4]?'Current':'Pre-delimitation'}</span></div></div>`}).join('');
  const groups=[...acLinks,...noAc].filter(l=>l[3]>0).map((l,k)=>{const names=l[4].split('|');const lim=60;
    const chip=n=>`<span class="chip${hitVillage&&n===hitVillage?' hit':''}">${esc(n)}</span>`;
    return `<div><h4>${esc(acName(l[1]))}${l[1]>=0?` <span class="pill">AC ${esc(acNo(l[1]))}</span>`:''} <span class="pill">${NF.format(names.length)} villages</span></h4>
      <div class="chips" data-all="${k}">${names.slice(0,lim).map(chip).join('')}${names.length>lim?`<button class="linkbtn" data-more="${k}">Show all ${NF.format(names.length)}</button>`:''}</div></div>`}).join('');
  const multi=acLinks.filter(l=>l[2]>=5).length>1;
  R.innerHTML=`<div class="resulthead"><div><div class="pin">${esc(p[0])}</div><div class="meta">${esc(p[1])} · ${esc(distOf(p)||'—')} · ${esc(stateOf(p)||'—')}</div></div>
    <div class="pills"><span class="pill ${multi?'warning':'primary'}">${acLinks.length} Vidhan Sabha${acLinks.length===1?'':'s'} found</span><span class="pill">${NF.format(nVil)} villages</span></div></div>
    <div class="path">${path}</div>
    ${cards?`<div class="grid">${cards}</div>`:'<div class="empty">No assembly constituency boundary covers this PIN code in the source file.</div>'}
    ${groups?`<div class="villages">${groups}</div>`:'<p class="extra">No village boundary falls inside this PIN code (typical for city PIN codes).</p>'}
    <div class="extra"><span>Division <b>${esc(p[2]||'—')}</b></span><span>Region <b>${esc(p[3]||'—')}</b></span><span>Postal circle <b>${esc(p[4]||'—')}</b></span></div>`;
  const allGroups=[...acLinks,...noAc].filter(l=>l[3]>0);
  R.querySelectorAll('[data-more]').forEach(b=>b.addEventListener('click',()=>{const l=allGroups[+b.dataset.more];const box=b.parentElement;box.innerHTML=l[4].split('|').map(n=>`<span class="chip${hitVillage&&n===hitVillage?' hit':''}">${esc(n)}</span>`).join('')}));
  try{history.replaceState(null,'','#'+p[0])}catch(e){}
}

function showVillages(q){
  if(!VN)buildVillageIndex();
  const s=q.toLowerCase(),pre=[],mid=[];
  for(let k=0;k<VN.length&&pre.length<150;k++){const v=VN[k];if(v.startsWith(s))pre.push(k);else if(mid.length<150&&v.includes(s))mid.push(k)}
  const seen=new Map();
  for(const k of pre.concat(mid)){const key=VNAME[k]+'|'+LINKS[VL[k]][0]+'|'+acName(LINKS[VL[k]][1])+'|'+acNo(LINKS[VL[k]][1]);if(seen.has(key))seen.get(key).n++;else seen.set(key,{k,n:1})}
  const hits=[...seen.values()].slice(0,150);
  const R=$('result');
  if(!hits.length){R.innerHTML=`<div class="empty">No village matching <b>${esc(q)}</b>. Try a shorter part of the name.</div>`;return}
  const hl=n=>{const j=n.toLowerCase().indexOf(s);return j<0?esc(n):esc(n.slice(0,j))+'<mark>'+esc(n.slice(j,j+s.length))+'</mark>'+esc(n.slice(j+s.length))};
  R.innerHTML=`<div class="resulthead"><div><div class="pin" style="font-size:26px">Villages matching “${esc(q)}”</div><div class="meta">Click a PIN code to open its full hierarchy.</div></div><div class="pills"><span class="pill info">${hits.length>=150?'150+':hits.length} matches</span></div></div>
  <div class="tablewrap" style="max-height:460px;border:1px solid var(--line);border-radius:14px"><table class="table"><thead><tr><th>Village</th><th>PIN</th><th>Post office</th><th>Vidhan Sabha</th><th class="num">AC No.</th><th>District</th><th>State</th></tr></thead><tbody>${hits.map(({k,n})=>{const l=LINKS[VL[k]],p=PINS[l[0]];
    return `<tr><td>${hl(VNAME[k])}${n>1?` <span class="pill">×${n}</span>`:''}</td><td><a href="#${esc(p[0])}" data-pin="${esc(p[0])}" data-vil="${esc(VNAME[k])}">${esc(p[0])}</a></td><td>${esc(p[1])}</td><td>${esc(acName(l[1]))}</td><td class="num">${esc(acNo(l[1]))}</td><td>${esc(distOf(p))}</td><td>${esc(stateOf(p))}</td></tr>`}).join('')}</tbody></table></div>`;
}

function search(){
  const q=$('q').value.trim();
  if(!q){$('result').innerHTML='<div class="empty">Enter a 6-digit PIN code or at least 3 letters of a village name.</div>';return}
  if(/^\d+$/.test(q)){if(!/^\d{6}$/.test(q)){$('result').innerHTML='<div class="empty">Please enter a valid 6-digit PIN code.</div>';return}showPin(q);return}
  if(q.length<3){$('result').innerHTML='<div class="empty">Type at least 3 letters of a village name.</div>';return}
  showVillages(q);
}
$('go').addEventListener('click',search);
$('q').addEventListener('keydown',e=>{if(e.key==='Enter')search()});
document.querySelector('.examples').addEventListener('click',e=>{const b=e.target.closest('[data-ex]');if(b){$('q').value=b.dataset.ex;search()}});
document.addEventListener('click',e=>{const a=e.target.closest('a[data-pin]');if(!a)return;e.preventDefault();$('q').value=a.dataset.pin;showPin(a.dataset.pin,a.dataset.vil||null);$('result').scrollIntoView({behavior:'smooth',block:'start'})});

// ---- explorer table
const sf=$('stateFilter'),df=$('distFilter'),ts=$('tableSearch');
[...new Set(PINS.map(p=>p[5]).filter(x=>x>=0))].map(x=>ST[x]).sort().forEach(s=>sf.add(new Option(s,s)));
let filtered=[];
function fillDistricts(){
  const s=sf.value;df.replaceChildren(new Option('All districts',''));df.disabled=!s;
  if(s)[...new Set(PINS.filter(p=>stateOf(p)===s).map(distOf).filter(Boolean))].sort().forEach(d=>df.add(new Option(d,d)));
}
function renderTable(){
  const s=sf.value,d=df.value,f=ts.value.trim().toLowerCase();
  filtered=SUM.filter(r=>{const p=PINS[r.i];if(s&&stateOf(p)!==s)return false;if(d&&distOf(p)!==d)return false;
    if(f){const hay=(p[0]+' '+p[1]+' '+distOf(p)+' '+stateOf(p)+' '+(r.ac>=0?acName(r.ac):'')).toLowerCase();if(!hay.includes(f))return false}return true});
  const LIM=500;
  $('tbody').innerHTML=filtered.slice(0,LIM).map(r=>{const p=PINS[r.i];return `<tr><td><a href="#${esc(p[0])}" data-pin="${esc(p[0])}">${esc(p[0])}</a></td><td>${esc(p[1])}</td><td class="num">${NF.format(p[7])}</td><td>${esc(r.ac>=0?acName(r.ac):'—')}</td><td class="num">${esc(r.ac>=0?acNo(r.ac):'—')}</td><td>${esc(distOf(p)||'—')}</td><td>${esc(stateOf(p)||'—')}</td><td class="num">${r.pct==null?'—':`<span class="mini"><i style="width:${Math.min(100,r.pct)}%"></i></span>${r.pct.toFixed(1)}%`}</td><td class="num">${r.n}</td></tr>`}).join('');
  $('count').textContent=`${NF.format(filtered.length)} PIN codes`;
  $('tablenote').textContent=filtered.length>LIM?`Showing the first ${LIM} of ${NF.format(filtered.length)}. Narrow with the filters, or download all matching rows as CSV.`:`Overlap = share of the PIN code's area in its primary Vidhan Sabha. ACs = constituencies covering at least 1% of the PIN code or holding one of its villages.`;
}
sf.addEventListener('change',()=>{fillDistricts();renderTable()});
df.addEventListener('change',renderTable);
let t=null;ts.addEventListener('input',()=>{clearTimeout(t);t=setTimeout(renderTable,150)});
$('csv').addEventListener('click',()=>{
  const q=v=>{v=String(v??'');return /[",\n]/.test(v)?'"'+v.replace(/"/g,'""')+'"':v};
  const rows=[['PIN','Post office','Villages','Vidhan Sabha','AC No.','Overlap %','Other Vidhan Sabhas','District','State','Division','Region','Circle']];
  filtered.forEach(r=>{const p=PINS[r.i];const ls=linksOf(r.i).filter(l=>l[1]>=0);
    rows.push([p[0],p[1],p[7],r.ac>=0?acName(r.ac):'',r.ac>=0?acNo(r.ac):'',r.pct==null?'':r.pct.toFixed(2),ls.slice(1).map(l=>`${acName(l[1])} (${acNo(l[1])}, ${l[2].toFixed(1)}%)`).join('; '),distOf(p),stateOf(p),p[2],p[3],p[4]])});
  const blob=new Blob(['﻿'+rows.map(r=>r.map(q).join(',')).join('\n')],{type:'text/csv;charset=utf-8'});
  const a=document.createElement('a');a.href=URL.createObjectURL(blob);a.download='pincode_vidhan_sabha'+(sf.value?'_'+sf.value.replace(/\W+/g,'_'):'')+'.csv';document.body.append(a);a.click();setTimeout(()=>{URL.revokeObjectURL(a.href);a.remove()},500);
});
fillDistricts();renderTable();
const h=decodeURIComponent(location.hash.slice(1));
if(/^\d{6}$/.test(h)){$('q').value=h;showPin(h)}
})();
</script>
</body>
</html>
'''

if __name__ == "__main__":
    main()
